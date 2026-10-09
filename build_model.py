"""
VYTA - Step 2: Build the data model
------------------------------------
Run with:   py build_model.py                  (uses the data in  data\\ )
       or:  py build_model.py data_stressed    (uses a different data folder, e.g. a test copy)

Reads the raw Xero files in  data\\raw\\  (made by pull_data.py) and creates:
  data\\model\\vyta_model.xlsx        - open in Excel to review the business month by month
  data\\model\\monthly_summary.csv    - the clean table score.py reads

Needs one extra package (one-off):   pip install openpyxl
Read-only: never touches Xero, only reads the files already on your computer.
"""

import csv
import json
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("\n[STOPPED] One package is missing. Run:  pip install openpyxl   then try again.")
    sys.exit(1)

RAW = Path("data") / "raw"
OUT = Path("data") / "model"

MONEY = '$#,##0;($#,##0);"-"'
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", start_color="1F3864")
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
CHECK_FILL = PatternFill("solid", start_color="F2F2F2")

# Liability lines treated as ATO / statutory debt
ATO_PATTERN = re.compile(r"gst|payg|superannuation|income tax|\bato\b|tax payable|\bbas\b", re.I)


# ---------- helpers ----------

def stop(message):
    print(f"\n[STOPPED] {message}")
    sys.exit(1)


def load(name):
    path = RAW / f"{name}.json"
    if not path.exists():
        stop(f"Can't find {path}. Run  py pull_data.py  first.")
    return json.loads(path.read_text())


def num(value):
    if value in (None, ""):
        return 0.0
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return 0.0


def period_end():
    summary = RAW / "summary.txt"
    if not summary.exists():
        stop(f"Can't find {summary}. Run  py pull_data.py  first.")
    match = re.search(r"ending (\d{4}-\d{2}-\d{2})", summary.read_text())
    if not match:
        stop("Couldn't read the period from summary.txt. Run  py pull_data.py  again.")
    return date.fromisoformat(match.group(1))


def month_ends(latest, count=12):
    """Newest first - matches the column order in Xero's reports."""
    ends = [latest]
    for _ in range(count - 1):
        ends.append(ends[-1].replace(day=1) - timedelta(days=1))
    return ends


def flatten(report_json, ends):
    """Turn Xero's nested report into simple rows: section, line, month, value."""
    rows = report_json["Reports"][0]["Rows"]
    out = []
    for block in rows:
        kind = block.get("RowType")
        if kind == "Section":
            section = block.get("Title", "") or ""
            children = block.get("Rows", [])
        elif kind in ("Row", "SummaryRow"):
            section, children = "", [block]
        else:
            continue
        for row in children:
            if row.get("RowType") not in ("Row", "SummaryRow"):
                continue
            cells = row.get("Cells", [])
            if not cells:
                continue
            label = (cells[0].get("Value") or "").strip()
            for i, cell in enumerate(cells[1:]):
                if i >= len(ends):
                    break
                out.append({
                    "section": section.strip(),
                    "line": label,
                    "is_total": row["RowType"] == "SummaryRow",
                    "month": ends[i],
                    "value": num(cell.get("Value")),
                })
    return out


def find(lines, month, line=None, section=None, total=None):
    """Value of the first line matching the patterns for that month (0 if not found)."""
    for r in lines:
        if r["month"] != month:
            continue
        if total is not None and r["is_total"] != total:
            continue
        if section and not re.search(section, r["section"], re.I):
            continue
        if line and not re.search(line, r["line"], re.I):
            continue
        return r["value"]
    return 0.0


def parse_xero_date(inv, key):
    text = inv.get(f"{key}String")
    if text:
        try:
            return datetime.fromisoformat(text[:19]).date()
        except ValueError:
            pass
    raw = inv.get(key) or ""
    match = re.search(r"/Date\((\d+)", raw)
    if match:
        return datetime.fromtimestamp(int(match.group(1)) / 1000, tz=timezone.utc).date()
    return None


# ---------- build ----------

def build_monthly(pl, bs, ends):
    ato_names = sorted({r["line"] for r in bs
                        if not r["is_total"] and re.search(r"liabilit", r["section"], re.I)
                        and ATO_PATTERN.search(r["line"])})
    months = []
    for m in sorted(ends):  # oldest first for reading left-to-right / top-to-bottom
        cash = find(bs, m, line=r"^total bank", section=r"^bank$", total=True)
        other_ca = find(bs, m, line=r"^total current assets", total=True)
        row = {
            "month": m,
            "revenue": find(pl, m, line=r"^total", section=r"^(?!.*other).*income", total=True),
            "cost_of_sales": find(pl, m, line=r"^total", section=r"cost of sales", total=True),
            "operating_expenses": find(pl, m, line=r"^total", section=r"operating expenses", total=True),
            "other_income": find(pl, m, line=r"^total", section=r"other income", total=True),
            "gross_profit": find(pl, m, line=r"^gross profit$"),
            "net_profit": find(pl, m, line=r"^net profit$"),
            "cash": cash,
            "debtors": find(bs, m, line=r"^accounts receivable$", total=False),
            "current_assets": cash + other_ca,
            "creditors": find(bs, m, line=r"^accounts payable$", total=False),
            "current_liabilities": find(bs, m, line=r"^total current liabilities$"),
            "total_assets": find(bs, m, line=r"^total assets$"),
            "total_liabilities": find(bs, m, line=r"^total liabilities$"),
            "net_assets": find(bs, m, line=r"^net assets$"),
        }
        for name in ato_names:
            row[f"ato: {name}"] = find(bs, m, line=f"^{re.escape(name)}$", total=False)
        row["ato_liabilities"] = sum(row[f"ato: {n}"] for n in ato_names)
        row["total_expenses"] = row["cost_of_sales"] + row["operating_expenses"]
        row["has_activity"] = any(abs(row[k]) > 0 for k in ("revenue", "operating_expenses", "cost_of_sales"))
        months.append(row)
    return months, ato_names


def build_aged(invoices, as_at):
    buckets = ["Not yet due", "1-30 days", "31-60 days", "61-90 days", "90+ days"]
    aged = {"ACCREC": {b: 0.0 for b in buckets}, "ACCPAY": {b: 0.0 for b in buckets}}
    overdue = []
    for inv in invoices:
        kind = inv.get("Type")
        due_amt = num(inv.get("AmountDue"))
        if kind not in aged or inv.get("Status") != "AUTHORISED" or due_amt <= 0:
            continue
        due = parse_xero_date(inv, "DueDate")
        days = (as_at - due).days if due else 0
        if days <= 0:
            bucket = buckets[0]
        elif days <= 30:
            bucket = buckets[1]
        elif days <= 60:
            bucket = buckets[2]
        elif days <= 90:
            bucket = buckets[3]
        else:
            bucket = buckets[4]
        aged[kind][bucket] += due_amt
        if days > 0:
            overdue.append({
                "type": "Owed to us" if kind == "ACCREC" else "We owe",
                "contact": (inv.get("Contact") or {}).get("Name", ""),
                "number": inv.get("InvoiceNumber", ""),
                "due": due,
                "days": days,
                "amount": due_amt,
            })
    overdue.sort(key=lambda x: (-x["days"], -x["amount"]))
    return buckets, aged, overdue


# ---------- Excel output ----------

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(path, org, end, months, ato_names, pl, bs, buckets, aged, overdue, as_at):
    wb = Workbook()

    # Sheet 1: Monthly summary
    ws = wb.active
    ws.title = "Monthly summary"
    columns = [
        ("Month", "month"), ("Revenue", "revenue"), ("Cost of sales", "cost_of_sales"),
        ("Operating expenses", "operating_expenses"), ("Total expenses", None),
        ("Other income", "other_income"), ("Gross profit", "gross_profit"),
        ("Net profit", "net_profit"), ("Cash", "cash"), ("Debtors", "debtors"),
        ("Current assets", "current_assets"), ("Creditors", "creditors"),
        ("Current liabilities", "current_liabilities"), ("Total assets", "total_assets"),
        ("Total liabilities", "total_liabilities"), ("Net assets", "net_assets"),
        ("ATO liabilities (total)", "ato_liabilities"),
    ] + [(f"ATO: {n}", f"ato: {n}") for n in ato_names] + [
        ("Check: P&L ties", None), ("Check: balance sheet ties", None), ("Has activity", "has_activity"),
    ]
    headers = [c[0] for c in columns]
    col = {h: get_column_letter(i) for i, h in enumerate(headers, start=1)}
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 32

    for r, m in enumerate(months, start=2):
        for c, (header, key) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT
            if header == "Month":
                cell.value = m["month"].strftime("%b %Y")
            elif header == "Total expenses":
                cell.value = f"={col['Cost of sales']}{r}+{col['Operating expenses']}{r}"
                cell.number_format = MONEY
            elif header == "Check: P&L ties":
                cell.value = (f"=ROUND({col['Revenue']}{r}-{col['Total expenses']}{r}"
                              f"+{col['Other income']}{r}-{col['Net profit']}{r},0)")
                cell.number_format = MONEY
                cell.fill = CHECK_FILL
            elif header == "Check: balance sheet ties":
                cell.value = (f"=ROUND({col['Total assets']}{r}-{col['Total liabilities']}{r}"
                              f"-{col['Net assets']}{r},0)")
                cell.number_format = MONEY
                cell.fill = CHECK_FILL
            elif key == "has_activity":
                cell.value = "Yes" if m["has_activity"] else "No"
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.value = round(m[key], 2)
                cell.number_format = MONEY
    ws.freeze_panes = "B2"
    set_widths(ws, [11] + [15] * (len(headers) - 1))

    # Sheets 2 & 3: line-by-line detail
    for title, lines in (("P&L detail", pl), ("Balance sheet detail", bs)):
        d = wb.create_sheet(title)
        d.append(["Month", "Section", "Line", "Type", "Value"])
        style_header(d, 1, 5)
        for r in sorted(lines, key=lambda x: x["month"]):
            d.append([r["month"].strftime("%b %Y"), r["section"], r["line"],
                      "Total" if r["is_total"] else "Account", round(r["value"], 2)])
        for row in d.iter_rows(min_row=2):
            for cell in row:
                cell.font = FONT
            row[4].number_format = MONEY
        d.freeze_panes = "A2"
        d.auto_filter.ref = d.dimensions
        set_widths(d, [11, 26, 36, 10, 15])

    # Sheet 4: aged snapshot + overdue list
    a = wb.create_sheet("Aged snapshot")
    a.append([f"Unpaid invoices as at {as_at:%d %b %Y}"])
    a["A1"].font = BOLD
    a.append([])
    a.append(["Bucket", "Owed to us (receivables)", "We owe (payables)"])
    style_header(a, 3, 3)
    for b in buckets:
        a.append([b, round(aged["ACCREC"][b], 2), round(aged["ACCPAY"][b], 2)])
    a.append(["Total", "=SUM(B4:B8)", "=SUM(C4:C8)"])
    for row in a.iter_rows(min_row=4, max_row=9):
        for cell in row:
            cell.font = BOLD if row[0].value == "Total" else FONT
        row[1].number_format = MONEY
        row[2].number_format = MONEY
    a.append([])
    a.append([])
    a.append(["Overdue invoices (most overdue first)"])
    a.cell(row=a.max_row, column=1).font = BOLD
    a.append(["Type", "Contact", "Invoice", "Due date", "Days overdue", "Amount due"])
    style_header(a, a.max_row, 6)
    for o in overdue:
        a.append([o["type"], o["contact"], o["number"],
                  o["due"].strftime("%d %b %Y") if o["due"] else "", o["days"], round(o["amount"], 2)])
        for cell in a[a.max_row]:
            cell.font = FONT
        a.cell(row=a.max_row, column=6).number_format = MONEY
    set_widths(a, [24, 30, 18, 14, 14, 15])

    # Sheet 5: notes
    n = wb.create_sheet("Notes")
    active = [m["month"].strftime("%b %Y") for m in months if m["has_activity"]]
    notes = [
        ("Vyta data model", ""),
        ("Organisation", org),
        ("Period", f"{len(months)} months ending {end:%d %b %Y}"),
        ("Built on", date.today().strftime("%d %b %Y")),
        ("Source", "Xero API (read-only) via pull_data.py - figures are hardcoded from Xero's own reports"),
        ("", ""),
        ("Months with activity", ", ".join(active) if active else "None found"),
        ("ATO lines matched", ", ".join(ato_names) if ato_names else "None found - check Balance sheet detail"),
        ("", ""),
        ("How to read it", ""),
        ("Current assets", "Bank balances + Xero's 'Total Current Assets' (Xero lists bank separately)"),
        ("ATO liabilities", "Sum of liability lines matching GST, PAYG, superannuation, income tax or ATO"),
        ("Check columns", "Should be $0 / '-'. Anything else means a line was missed - tell Claude"),
        ("Aged snapshot", "Uses today's unpaid balances (not month-end); approved invoices only"),
        ("Grey cells", "Formulas - they recalculate if you change a number"),
    ]
    for label, value in notes:
        n.append([label, value])
    for row in n.iter_rows():
        row[0].font = BOLD
        row[1].font = FONT
    n["A1"].font = Font(name="Arial", size=12, bold=True)
    set_widths(n, [22, 95])

    wb.move_sheet("Notes", offset=-(len(wb.sheetnames) - 1))
    wb.active = 1
    wb.save(path)


def write_csv(path, months, ato_names):
    keys = ["month", "revenue", "cost_of_sales", "operating_expenses", "total_expenses",
            "other_income", "gross_profit", "net_profit", "cash", "debtors", "current_assets",
            "creditors", "current_liabilities", "total_assets", "total_liabilities",
            "net_assets", "ato_liabilities"] + [f"ato: {n}" for n in ato_names] + ["has_activity"]
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(keys)
        for m in months:
            w.writerow([m["month"].isoformat() if k == "month"
                        else ("yes" if m[k] else "no") if k == "has_activity"
                        else round(m[k], 2) for k in keys])


def main():
    global RAW, OUT
    base = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data")
    RAW, OUT = base / "raw", base / "model"

    end = period_end()
    pl_report = load("profit_and_loss")
    header = next((r for r in pl_report["Reports"][0]["Rows"] if r.get("RowType") == "Header"), None)
    ends = month_ends(end, len(header["Cells"]) - 1 if header else 12)
    org_data = load("organisation")
    org = (org_data.get("Organisations") or [{}])[0].get("Name", "Unknown")

    pl = flatten(pl_report, ends)
    bs = flatten(load("balance_sheet"), ends)
    invoices = load("invoices")

    months, ato_names = build_monthly(pl, bs, ends)
    as_at = date.today()
    buckets, aged, overdue = build_aged(invoices, as_at)

    OUT.mkdir(parents=True, exist_ok=True)
    xlsx_path = OUT / "vyta_model.xlsx"
    csv_path = OUT / "monthly_summary.csv"
    try:
        write_workbook(xlsx_path, org, end, months, ato_names, pl, bs, buckets, aged, overdue, as_at)
    except PermissionError:
        stop("vyta_model.xlsx is open in Excel. Close it and run again.")
    write_csv(csv_path, months, ato_names)

    # Terminal summary + self-checks
    print(f"Organisation: {org}")
    print(f"Period:       {len(months)} months ending {end:%d %b %Y}\n")
    active = [m["month"].strftime("%b %Y") for m in months if m["has_activity"]]
    print(f"Months with activity ({len(active)}/{len(months)}): {', '.join(active) if active else 'none'}")
    print(f"ATO lines found: {', '.join(ato_names) if ato_names else 'none'}")

    pl_bad = [m["month"].strftime("%b %Y") for m in months
              if abs(m["revenue"] - m["total_expenses"] + m["other_income"] - m["net_profit"]) > 1]
    bs_bad = [m["month"].strftime("%b %Y") for m in months
              if abs(m["total_assets"] - m["total_liabilities"] - m["net_assets"]) > 1]
    print(f"P&L ties:           {'OK' if not pl_bad else 'MISMATCH in ' + ', '.join(pl_bad)}")
    print(f"Balance sheet ties: {'OK' if not bs_bad else 'MISMATCH in ' + ', '.join(bs_bad)}")

    latest = months[-1]
    print(f"\nLatest month ({latest['month']:%b %Y}): revenue ${latest['revenue']:,.0f}, "
          f"net profit ${latest['net_profit']:,.0f}, cash ${latest['cash']:,.0f}, "
          f"ATO liabilities ${latest['ato_liabilities']:,.0f}")
    print(f"Overdue invoices: {sum(1 for o in overdue if o['type'] == 'Owed to us')} owed to us, "
          f"{sum(1 for o in overdue if o['type'] == 'We owe')} we owe")

    print(f"\n[DONE] Saved:\n  {xlsx_path.resolve()}\n  {csv_path.resolve()}")


if __name__ == "__main__":
    main()
