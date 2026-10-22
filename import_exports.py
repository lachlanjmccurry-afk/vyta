"""
VYTA - Import Xero Excel exports for testing
---------------------------------------------
Run with:   py import_exports.py            (every business in "Data Exports")
            py import_exports.py --list     (just show what it found)

Turns Xero report exports (Excel) into the same raw files pull_data.py gets from the Xero API, so the
normal pipeline scores them unchanged. Each business gets its own folder, data_business_a, data_business_b
and so on (ignored by git, like every data_* folder), then build_model.py and score.py are run on it.

What it reads, from the "Data Exports" folder:
  - each workbook in the folder is one business (e.g. an investigation workbook with several sheets)
  - each subfolder is one business (e.g. "Xero Extracts", one report per file)
  Sheets it uses: Profit and Loss (monthly), Balance Sheet (monthly), Aged Receivables / Payables
  (Detail if there is one, otherwise Summary). Yearly, "(Adjusted)", ratio and ledger sheets are skipped.

Privacy:
  - Businesses are named "Business A", "Business B"... in everything Vyta shows; the real name is never copied.
  - Customer and supplier names from aged reports are replaced with "Customer 1", "Supplier 1"...
  - The mapping from letter to source file is printed here only (not saved anywhere).
  - Never commit data_business_* folders or "Data Exports" (both are in .gitignore).

How the reports are reshaped (Xero layouts differ from business to business):
  - P&L sections are mapped to Vyta's four: Trading Income, Cost of Sales, Other Income, Operating Expenses.
    Every expense section below Gross Profit (Marketing, Wages, Interest and Depreciation...) counts as
    Operating Expenses. Lines keep their names, so V5's interest and depreciation add-backs still work.
  - Only months that appear in BOTH the P&L and the balance sheet are used.
  - If the four P&L sections don't add up to the reported Net Profit, Operating Expenses is set to the
    balancing figure so the P&L ties, and the difference is reported.
  - Non-current liability lines are marked as term loans (Xero's TERMLIAB) for the bank-debt checks.
  - Aged invoices: Detail rows keep their real due dates. Summary buckets get a representative due date
    (< 1 Month: current, 1 Month: 45 days overdue, 2 Months: 75, 3 Months: 105, Older: 150 days), which is
    exact for Vyta's "more than 90 days overdue" tests.
  - If a business has no aged receivables or payables report, D8 / D10 show "Not assessed" instead of Green.
"""

import calendar
import json
import re
import subprocess
import sys
import warnings
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("Data Exports")
LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_name) if m})
MONTHS["sept"] = 9
BUCKET_DAYS = {"< 1 month": 0, "1 month": 45, "2 months": 75, "3 months": 105, "older": 150}

warnings.filterwarnings("ignore")  # openpyxl warns about Excel data validation; harmless


def stop(message):
    print(f"[STOPPED] {message}")
    sys.exit(1)


# ---------- reading sheets ----------

def parse_month(text):
    """'Sept 2026', '28 Sept 2026', 'June 2026' -> (2026, 9). None if it isn't a month heading."""
    if not isinstance(text, str):
        return None
    m = re.fullmatch(r"\s*(?:\d{1,2}\s+)?([A-Za-z]+)\s+(\d{4})\s*", text)
    if not m or m.group(1).lower() not in MONTHS:
        return None
    return int(m.group(2)), MONTHS[m.group(1).lower()]


def month_end(ym):
    y, m = ym
    return date(y, m, calendar.monthrange(y, m)[1])


def as_at_date(rows):
    """The 'As at 31 October 2026' line in a report's title rows."""
    for r in rows[:6]:
        for c in r:
            if isinstance(c, str):
                m = re.search(r"as at\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", c, re.I)
                if m and m.group(2).lower() in MONTHS:
                    return date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1)))
    return None


def statement_rows(rows):
    """Read a Xero P&L or balance sheet export: [(level, label, {(y, m): value} or None)], month list."""
    head = next((i for i, r in enumerate(rows[:15]) if any(c == "Account" for c in r)), None)
    if head is None:
        return [], []
    cols = {j: parse_month(c) for j, c in enumerate(rows[head]) if parse_month(c)}
    if not cols:
        return [], []
    first_col = min(cols)
    out = []
    for r in rows[head + 1:]:
        level, label = next(((j, c.strip()) for j, c in enumerate(r[:first_col]) if isinstance(c, str) and c.strip()),
                            (None, None))
        if label is None:
            continue
        if label.startswith("[FX]"):
            break  # footnote about exchange rates: the statement has ended
        nums = {ym: r[j] for j, ym in cols.items() if j < len(r) and isinstance(r[j], (int, float))}
        values = {ym: float(nums.get(ym, 0.0)) for ym in cols.values()} if nums else None
        out.append((level, label, values))
    return out, list(cols.values())


def is_monthly(rows):
    _, months = statement_rows(rows)
    return len(months) >= 6 and len(set(months)) == len(months)


def pick_sheet(sheets, pattern):
    """The monthly version of a report: skips Yearly / Adjusted / No Loan variants."""
    for title, rows in sheets:
        if re.match(pattern, title, re.I) and not re.search(r"yearly|adjusted|no loan", title, re.I) and is_monthly(rows):
            return rows
    return None


# ---------- reshaping into Xero API report JSON ----------

def cell(v):
    return {"Value": f"{v:.2f}" if isinstance(v, float) else v}


def report_json(name, months, sections):
    """sections: [(title, [(row_type, label, {ym: value})])] -> the shape build_model.flatten() reads."""
    header = {"RowType": "Header", "Cells": [cell("")] + [cell(f"{month_end(ym):%d %b %Y}") for ym in months]}
    rows = [header]
    for title, lines in sections:
        rows.append({"RowType": "Section", "Title": title, "Rows": [
            {"RowType": kind, "Cells": [cell(label)] + [cell(vals.get(ym, 0.0)) for ym in months]}
            for kind, label, vals in lines]})
    return {"Reports": [{"ReportName": name, "Rows": rows}]}


def pl_sections(rows, months):
    """Map a Xero P&L export onto Vyta's four sections. Returns (sections, gross, net, warnings)."""
    role_of = lambda title, after_gp: (  # order matters: "Cost of Sales" also contains "sales"
        "Other Income" if re.search(r"other income", title, re.I) else
        "Cost of Sales" if re.search(r"cost of (sales|goods)|direct cost|purchases", title, re.I) else
        "Trading Income" if re.search(r"income|revenue|sales", title, re.I) and not after_gp else
        "Cost of Sales" if not after_gp else
        "Operating Expenses")
    lines = {k: [] for k in ("Trading Income", "Cost of Sales", "Other Income", "Operating Expenses")}
    groups = []  # one per section in the export: {"role", "title", "lines", "reported" (Xero's total row)}
    gross = net = None
    after_gp = False
    for level, label, vals in rows:
        low = label.lower()
        if re.match(r"gross profit", low):
            gross, after_gp = vals, True
            continue
        if re.match(r"net profit", low):
            net = vals
            break  # anything after Net Profit (e.g. tax, distributions) isn't part of operating profit
        if vals is None:
            if level == 0:  # a top-level heading starts a new section
                groups.append({"role": role_of(label, after_gp), "title": label, "lines": [], "reported": None})
            continue
        if not groups:  # lines before any heading (seen in compact layouts)
            groups.append({"role": role_of(label, after_gp), "title": label, "lines": [], "reported": None})
        if low.startswith("total "):
            if level == 0:
                groups[-1]["reported"] = vals
            continue  # sub-totals inside a section are left out so nothing is counted twice
        groups[-1]["lines"].append(("Row", label, vals))
        lines[groups[-1]["role"]].append(("Row", label, vals))
    if net is None:
        return None, None, None, ["no Net Profit line"]

    # Each role's total: Xero's own section total where it printed one, otherwise the sum of the lines
    tot = {k: {ym: 0.0 for ym in months} for k in lines}
    notes = []
    for g in groups:
        line_sum = {ym: sum(v.get(ym, 0.0) for _, _, v in g["lines"]) for ym in months}
        use = line_sum
        if g["reported"] is not None:
            off = max(abs(g["reported"].get(ym, 0.0) - line_sum[ym]) for ym in months)
            if off > 1:
                notes.append(f"{g['title']}: lines differ from Xero's total by up to ${off:,.0f} (used Xero's total)")
            use = {ym: g["reported"].get(ym, 0.0) for ym in months}
        for ym in months:
            tot[g["role"]][ym] += use[ym]
    gp = {ym: tot["Trading Income"][ym] - tot["Cost of Sales"][ym] for ym in months}
    if gross is not None:
        off = max(abs(gp[ym] - gross.get(ym, 0.0)) for ym in months)
        if off > 1:
            notes.append(f"gross profit differs from Xero's by up to ${off:,.0f} (used Xero's)")
            gp = {ym: gross.get(ym, 0.0) for ym in months}
            tot["Cost of Sales"] = {ym: tot["Trading Income"][ym] - gp[ym] for ym in months}
    balancing = {ym: gp[ym] + tot["Other Income"][ym] - net.get(ym, 0.0) for ym in months}
    off = max(abs(balancing[ym] - tot["Operating Expenses"][ym]) for ym in months)
    if off > 1:
        notes.append(f"operating expenses set to the balancing figure (differed by up to ${off:,.0f})")
        tot["Operating Expenses"] = balancing

    sections = []
    for key in ("Trading Income", "Cost of Sales"):
        sections.append((key, lines[key] + [("SummaryRow", f"Total {key}", tot[key])]))
    sections.append(("", [("Row", "Gross Profit", gp)]))
    for key in ("Other Income", "Operating Expenses"):
        sections.append((key, lines[key] + [("SummaryRow", f"Total {key}", tot[key])]))
    sections.append(("", [("Row", "Net Profit", {ym: net.get(ym, 0.0) for ym in months})]))
    return sections, gp, net, notes


def bs_sections(rows):
    """Keep the balance sheet as Xero laid it out; liability-side sub-sections are tagged '(Liabilities)'."""
    sections, top, title, current = [], "", "", []
    accounts = []

    def flush():
        if current:
            sections.append((title, list(current)))
            current.clear()

    for level, label, vals in rows:
        if vals is None:
            flush()
            if level == 0:
                top, title = label, label
            else:
                title = label if re.search(r"liabilit", label, re.I) or not re.search(r"liabilit", top, re.I) \
                    else f"{label} (Liabilities)"
            continue
        kind = "SummaryRow" if label.lower().startswith("total ") else "Row"
        current.append((kind, label, vals))
        if kind == "Row" and re.search(r"non-?current liabilit", title, re.I):
            accounts.append({"Name": label, "Type": "TERMLIAB"})
    flush()
    return sections, accounts


def aged_invoices(rows, kind, as_at, names):
    """Invoices from an Aged Receivables/Payables export (Detail if it has due dates, else Summary buckets)."""
    who = "Customer" if kind == "ACCREC" else "Supplier"
    head = next((i for i, r in enumerate(rows[:10]) if any(isinstance(c, str) and c.strip() in ("Total",) for c in r)), None)
    if head is None:
        return []
    hdr = [c.strip().lower() if isinstance(c, str) else "" for c in rows[head]]
    if "total" not in hdr:
        return []
    t_col = hdr.index("total")
    out = []

    def contact(name):
        return names.setdefault((kind, name), f"{who} {sum(1 for k in names if k[0] == kind) + 1}")

    if "due date" in hdr:  # Detail: one row per invoice under a contact heading
        due_col = hdr.index("due date")
        current = None
        for r in rows[head + 1:]:
            first = r[0]
            if isinstance(first, str) and not any(isinstance(c, (int, float)) for c in r):
                current = first.strip()
                continue
            due = r[due_col] if due_col < len(r) else None
            amount = r[t_col] if t_col < len(r) else None
            if isinstance(due, datetime) and isinstance(amount, (int, float)) and amount > 0:
                out.append({"Type": kind, "Status": "AUTHORISED", "Contact": {"Name": contact(current)},
                            "AmountDue": round(float(amount), 2),
                            "DueDateString": due.strftime("%Y-%m-%dT00:00:00")})
    else:  # Summary: one row per contact with ageing buckets
        buckets = {j: BUCKET_DAYS[h] for j, h in enumerate(hdr) if h in BUCKET_DAYS}
        for r in rows[head + 1:]:
            name = r[0]
            if not isinstance(name, str) or name.strip().lower().startswith(("total", "percentage")):
                continue
            for j, days in buckets.items():
                amount = r[j] if j < len(r) else None
                if isinstance(amount, (int, float)) and amount > 0:
                    due = as_at - timedelta(days=days)
                    out.append({"Type": kind, "Status": "AUTHORISED", "Contact": {"Name": contact(name.strip())},
                                "AmountDue": round(float(amount), 2),
                                "DueDateString": due.strftime("%Y-%m-%dT00:00:00")})
    return out


# ---------- one business ----------

def load_sheets(paths):
    sheets = []
    for p in paths:
        wb = load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            sheets.append((ws.title, list(ws.iter_rows(values_only=True))))
        wb.close()
    return sheets


def import_business(letter, paths):
    sheets = load_sheets(paths)
    pl_rows = pick_sheet(sheets, r"profit and loss")
    bs_rows = pick_sheet(sheets, r"balance sheet")
    if pl_rows is None or bs_rows is None:
        print(f"  Business {letter}: skipped (needs a monthly Profit and Loss and Balance Sheet)")
        return None

    pl, pl_months = statement_rows(pl_rows)
    bs, bs_months = statement_rows(bs_rows)
    months = sorted(set(pl_months) & set(bs_months), reverse=True)  # newest first, like Xero
    if len(months) < 3:
        print(f"  Business {letter}: skipped (P&L and balance sheet share only {len(months)} months)")
        return None

    pl_secs, gp, net, notes = pl_sections(pl, months)
    if pl_secs is None:
        print(f"  Business {letter}: skipped ({notes[0]})")
        return None
    bs_secs, accounts = bs_sections(bs)

    # Aged reports: prefer Detail over Summary for each side
    names, invoices, as_at, missing = {}, [], None, []
    for kind, word in (("ACCREC", "receivables"), ("ACCPAY", "payables")):
        found = None
        for style in ("detail", "summary"):
            for title, rows in sheets:
                if re.match(rf"aged {word} {style}", title, re.I):
                    when = as_at_date(rows) or month_end(months[0])
                    got = aged_invoices(rows, kind, when, names)
                    if got or (found is None and len(rows) > 6):
                        found = (got, when)
                        break
            if found and found[0]:
                break
        if found is None:
            missing.append(kind)
        else:
            invoices += found[0]
            as_at = max(as_at or found[1], found[1])

    end = month_end(months[0])
    as_at = as_at or end
    out = Path(f"data_business_{letter.lower()}")
    raw = out / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    (raw / "profit_and_loss.json").write_text(json.dumps(report_json("Profit and Loss", months, pl_secs), indent=1))
    (raw / "balance_sheet.json").write_text(json.dumps(report_json("Balance Sheet", months, bs_secs), indent=1))
    (raw / "invoices.json").write_text(json.dumps(invoices, indent=1))
    (raw / "accounts.json").write_text(json.dumps({"Accounts": accounts}, indent=1))
    (raw / "organisation.json").write_text(json.dumps({"Organisations": [{"Name": f"Business {letter}"}]}, indent=1))
    (raw / "aged_status.json").write_text(json.dumps({"missing": missing}, indent=1))
    (raw / "summary.txt").write_text(
        f"Organisation: Business {letter}\nPulled on:    {as_at.isoformat()}\n"
        f"Period:       {len(months)} months ending {end.isoformat()}\n\n"
        "IMPORTED FROM XERO EXCEL EXPORTS by import_exports.py (anonymised). Not pulled from the Xero API.\n"
        + "".join(f"Note: {n}\n" for n in notes)
        + "".join(f"Note: no aged {'receivables' if k == 'ACCREC' else 'payables'} report, so "
                  f"{'D10' if k == 'ACCREC' else 'D8'} is not assessed\n" for k in missing))

    print(f"  Business {letter}: {len(months)} months to {end:%b %Y}, {len(invoices)} unpaid invoices "
          f"(as at {as_at:%d %b %Y})" + (f", no aged {' or '.join('receivables' if k == 'ACCREC' else 'payables' for k in missing)}" if missing else ""))
    for n in notes:
        print(f"      note: {n}")
    return out


def businesses():
    """Each workbook in Data Exports is one business; each subfolder is one business."""
    if not SOURCE.exists():
        stop(f'Can\'t find the "{SOURCE}" folder. Make it in the vyta folder and put the Excel exports in it.')
    found = [[p] for p in sorted(SOURCE.glob("*.xls*")) if not p.name.startswith("~$")]
    found += [sorted(d.glob("*.xls*")) for d in sorted(SOURCE.iterdir()) if d.is_dir() and list(d.glob("*.xls*"))]
    return found


def main(args):
    groups = businesses()
    print(f"Found {len(groups)} business(es) in \"{SOURCE}\" (shown here only; Vyta uses the letters):")
    for letter, paths in zip(LETTERS, groups):
        print(f"  Business {letter} = {paths[0].parent.name + '/' if len(paths) > 1 else ''}{paths[0].name if len(paths) == 1 else f'{len(paths)} files'}")
    if "--list" in args:
        return
    print("\nImporting:")
    made = [f for f in (import_business(letter, paths) for letter, paths in zip(LETTERS, groups)) if f]
    print("\nBuilding and scoring:")
    for folder in made:
        for script in ("build_model.py", "score.py"):
            res = subprocess.run([sys.executable, script, str(folder)], capture_output=True, text=True)
            if res.returncode != 0:
                print(f"  {folder}: {script} stopped:\n{res.stdout[-800:]}{res.stderr[-800:]}")
                break
        else:
            print(f"  {folder}: done")
    print(f"\n[DONE] Imported {len(made)} business(es). They appear in the dashboard's business picker "
          "(py -m streamlit run dashboard.py).")


if __name__ == "__main__":
    main(sys.argv[1:])
